"""
Validación Factura Global — Consolidation: Monivoi (ZLMXCOM_TRN_MONINVOI)
=========================================================================

Everything that is computed FROM the Monivoi (the invoicing monitor), plus its
cross against the accounting side:

    1) 'MTD' sheet (step 16): four 'Centro' × 'Valor neto' matrices (SG, RG,
       PP and TODOS)                    → preparar_matrices_mtd / escribir_hoja_mtd
    2) Helper column 'Referencia 2' as column A of the 'Monivoi' sheet
                                        → insertar_referencia2
    3) TD × MTD cross: live Excel formulas written on the 'TD' sheet
       (Factura Monitor / Diferencia / Observación), plus the warnings those
       formulas cannot show on their own
                                        → escribir_cruce_td_mtd / avisos_cruce_td_mtd

The cross lives here because it is the Monivoi side of the reconciliation: its
formulas look up MTD values from the TD rows built by Consolidacion_Gallo. That
is why this module imports the TD layout constants from Consolidacion_Gallo
(one-way dependency: Gallo never imports Monivoi).

Dependencies: Consolidacion_Comun and Consolidacion_Gallo (TD layout only).
"""

import re
from functools import partial

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from Consolidacion_Comun import (
    ETIQUETA_GRAND_TOTAL,
    TD_FORMATO_NUMERO,
    col_por_candidatos,
    escribir_filtros,
    parse_importe_us,
)
from Consolidacion_Gallo import (
    TD_FACTURAS,
    TD_FILA_ENCABEZADO,
    TD_FILA_INICIO,
    TD_NOTAS,
    TD_PEP,
)

# ----------------------------------------------------------------------
# Process constants
# ----------------------------------------------------------------------
# Sheet names (tabs) built from the Monivoi.
HOJA_MONIVOI    = "Monivoi"
HOJA_MTD        = "MTD"

# --- 'MTD' sheet (step 16) ---
# Four 'Centro' × sum of 'Valor neto' matrices, all taken from the MONIVOI.
# Layout: rows 1-3 = filter legend, row 4 = title, row 5 = headers,
# row 6 onwards = data. The three top rows exist ONLY for the legend, so the
# stakeholder reads under which conditions the matrix was built before reading
# the numbers. The TD × MTD cross is NOT affected by these rows: it looks the
# centros up with XLOOKUP over WHOLE columns (MTD!A:A, MTD!B:B), never over a
# fixed range.
MTD_FILA_FILTROS    = 1
MTD_FILA_TITULO     = 4
MTD_FILA_ENCABEZADO = 5
MTD_FILA_INICIO     = 6

MTD_ENCABEZADO_CENTRO   = "Centro"
MTD_ENCABEZADO_VALOR    = "Valor neto"
MTD_ETIQUETA_TOTAL      = ETIQUETA_GRAND_TOTAL   # same as 'TD' (see the cross)
MTD_ETIQUETA_SIN_CENTRO = "Sin centro"

MTD_PREFIJO_SG = "SG"
MTD_PREFIJO_RG = "RG"
MTD_PREFIJO_PP = "PP"
# The three prefixes that go into the fourth matrix ('TODOS'). They also define
# the universe of rows that get parsed: nothing outside these three is touched.
MTD_PREFIJOS_TODOS = frozenset({MTD_PREFIJO_SG, MTD_PREFIJO_RG, MTD_PREFIJO_PP})

# (internal key, title, prefixes it filters, 'Centro' column, value column)
MTD_BLOQUES = (
    {"clave": "sg",    "titulo": "SG",
     "prefijos": frozenset({MTD_PREFIJO_SG}),
     "col_centro": 1,  "col_valor": 2},     # A / B
    {"clave": "rg",    "titulo": "RG",
     "prefijos": frozenset({MTD_PREFIJO_RG}),
     "col_centro": 4,  "col_valor": 5},     # D / E
    {"clave": "pp",    "titulo": "PP",
     "prefijos": frozenset({MTD_PREFIJO_PP}),
     "col_centro": 11, "col_valor": 12},    # K / L
    {"clave": "todos", "titulo": "TODOS (SG + RG + PP)",
     "prefijos": MTD_PREFIJOS_TODOS,
     "col_centro": 17, "col_valor": 18},    # Q / R
)

# Name of the helper column added AS THE FIRST COLUMN (A) of the 'Monivoi'
# sheet, so the matrices can be audited by filtering by hand.
MTD_COL_REFERENCIA2 = "Referencia 2"

# Monivoi columns that 'MTD' needs. Spelling variants are accepted (the
# comparison ignores accents, dots, spaces and case).
MTD_COL_REFERENCIA_CANDIDATAS = ("Referencia", "Referencia de factura")
MTD_COL_UUID_CANDIDATAS       = ("UUID", "UUID CFDI", "Folio fiscal", "UUID Fiscal")
MTD_COL_CENTRO_CANDIDATAS     = ("Centro", "Centro de distribución",
                                 "Puesto de expedición")
MTD_COL_VALOR_CANDIDATAS      = ("Valor neto", "Valor neto factura", "Importe neto")

# Texts that count as an EMPTY UUID (besides a blank or whitespace-only cell).
# They are compared upper-cased and stripped. Note: the SHAPE of the UUID is
# NOT validated here on purpose — if SAP exported it without dashes it would
# still be a real UUID, and we do not want to drop good invoices over a format
# change. Only the "there is nothing here" sentinels are discarded.
MTD_UUID_VACIOS = frozenset({
    "", "NULL", "NONE", "NAN", "NA", "N/A", "#N/A", "#N/D", "-", "--", "#",
})

# Tolerance of the 'TODOS == SG + RG + PP' check. One cent covers the rounding
# accumulated by adding values already rounded to 2 decimals.
MTD_TOLERANCIA_CUADRE = 0.01

# Direct access to each 'MTD' block by its key ('sg', 'rg', 'pp', 'todos').
_MTD_POR_CLAVE = {bloque["clave"]: bloque for bloque in MTD_BLOQUES}

# --- 'TD' × 'MTD' cross ---
# Three new columns to the right of each 'TD' block, with LIVE Excel FORMULAS
# (not values), so the stakeholder can audit them and drag them down:
#
#     Facturas     (A..D) ↔ MTD 'SG' (A/B) → E 'Factura Monitor', F 'Diferencia', G 'Observación'
#     Notas        (L..O) ↔ MTD 'RG' (D/E) → P, Q, R
#     PEP vouchers (W..X) ↔ MTD 'PP' (K/L) → Y, Z, AA
#
# The formula goes in EVERY centro row AND in the block's 'Grand Total' row:
# there the XLOOKUP finds the MTD 'Grand Total' and compares total against
# total, including centros that only exist in the monitor and the 'Sin centro'
# row (that is why TD and MTD share ETIQUETA_GRAND_TOTAL).
TD_ENCABEZADO_MONITOR     = "Factura Monitor"
TD_ENCABEZADO_DIFERENCIA  = "Diferencia"
TD_ENCABEZADO_OBSERVACION = "Observación"
TD_TEXTO_SIN_TIMBRE       = "Sin timbre"
TD_TEXTO_SIN_DIFERENCIA   = "Sin diferencia"

# 'Observación' tolerance: up to one cent counts as 'Sin diferencia'.
TD_CRUCE_TOLERANCIA = 0.01

# 'revisar_si' marks the ONLY direction that leaves the Observación blank for
# review: when accounting exceeds the monitor. The opposite direction shows
# 'Sin diferencia' ON PURPOSE (confirmed business rule).
#     '<' → the Diferencia is compared against -tolerance  (ROUND(F6,2)<-0.01)
#     '>' → the Diferencia is compared against +tolerance  (ROUND(Q6,2)>0.01)
TD_CRUCES = (
    {   # Facturas ↔ SG. Diferencia = D + E: invoice accounting comes in as
        # negative, so a perfect match gives zero.
        "etiqueta":        "Facturas ↔ SG",
        "td_clave":        "facturas",
        "col_cedis":       TD_FACTURAS["col_cedis"],        # A
        "col_contab":      TD_FACTURAS["col_grand_total"],  # D
        "col_monitor":     5,                               # E
        "col_diferencia":  6,                               # F
        "col_observacion": 7,                               # G
        "mtd_clave":       "sg",                            # MTD A / B
        "operador":        "+",
        "revisar_si":      "<",
    },
    {   # Notas de Crédito ↔ RG. Diferencia = O - P.
        "etiqueta":        "Notas de Crédito ↔ RG",
        "td_clave":        "notas",
        "col_cedis":       TD_NOTAS["col_cedis"],           # L
        "col_contab":      TD_NOTAS["col_grand_total"],     # O
        "col_monitor":     16,                              # P
        "col_diferencia":  17,                              # Q
        "col_observacion": 18,                              # R
        "mtd_clave":       "rg",                            # MTD D / E
        "operador":        "-",
        "revisar_si":      ">",
    },
    {   # PEP vouchers ↔ PP. Diferencia = X - Y.
        "etiqueta":        "PEP vouchers ↔ PP",
        "td_clave":        "pep",
        "col_cedis":       TD_PEP["col_cedis"],             # W
        "col_contab":      TD_PEP["col_valor"],             # X
        "col_monitor":     25,                              # Y
        "col_diferencia":  26,                              # Z
        "col_observacion": 27,                              # AA
        "mtd_clave":       "pp",                            # MTD K / L
        "operador":        "-",
        "revisar_si":      ">",
    },
)

# A sheet name made ONLY of letters (and '_') can go unquoted in a formula
# ('MTD!A:A'); any other ('Venta DSD', 'Hoja1') needs single quotes. So do R,
# C, TRUE and FALSE, because Excel mistakes them for something else.
_RE_HOJA_SIN_COMILLAS = re.compile(r"[A-Za-z_]+")
_HOJAS_RESERVADAS = frozenset({"R", "C", "TRUE", "FALSE"})


# ----------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------
def _nombre_columna_libre(df, base):
    """
    Return `base` if that name does NOT exist yet in the DataFrame; otherwise
    return 'base (2)', 'base (3)', ...

    Used to add the helper column 'Referencia 2' to the Monivoi sheet without
    overwriting a same-named column the export may bring (overwriting it would
    silently wipe the stakeholder's original data).
    """
    if base not in df.columns:
        return base
    sufijo = 2
    while f"{base} ({sufijo})" in df.columns:
        sufijo += 1
    return f"{base} ({sufijo})"


def _uuid_vacio(valor):
    """
    True if the row's UUID counts as EMPTY or NULL (step 16).

    Covers the blank cell, the whitespace-only cell and the text sentinels the
    export sometimes leaves ('NULL', 'N/A', '-', ...). It does NOT validate the
    UUID's shape: a real UUID written differently (no dashes, lower case) still
    counts as present.
    """
    return str(valor).strip().upper() in MTD_UUID_VACIOS


# ----------------------------------------------------------------------
# 'MTD' sheet (step 16) — computation and writing
# ----------------------------------------------------------------------
def _mtd_bloques_vacios():
    """Empty structure (one per block) for when the Monivoi has no data."""
    return {
        bloque["clave"]: {
            "filas":              [],
            "sin_centro_n":       0,
            "sin_centro_importe": 0.0,
            "total":              0.0,
        }
        for bloque in MTD_BLOQUES
    }


def preparar_matrices_mtd(monivoi, monivoi_sin_datos):
    """
    Step 16 — build the FOUR matrices of the 'MTD' sheet, all of them over the
    MONIVOI sheet (not over the Gallo).

    Grouping key: the first 2 characters of the Monivoi 'Referencia' column.
    Filter shared by all four: the row's UUID can NOT be empty or null (see
    `_uuid_vacio`); those rows were never invoiced and including them would
    inflate the total against the monitor.

        Matrix 1 (A/B) → Referencia 2 == 'SG'
        Matrix 2 (D/E) → Referencia 2 == 'RG'
        Matrix 3 (K/L) → Referencia 2 == 'PP'
        Matrix 4 (Q/R) → Referencia 2 ∈ {'SG', 'RG', 'PP'}  (TODOS)

    Each matrix derives its OWN list of centros from ITS own filters, same as
    in 'TD', so the four blocks can have different centros and a different
    number of rows.

    'Valor neto' is only parsed for the rows that go into some matrix (the
    union of the three prefixes), not for the whole sheet.

    Rows that pass the filter but come WITHOUT a centro are not dropped: they
    are shown apart in a 'Sin centro' row, so the 'Grand Total' of each block
    always reconciles with its filter.

    Returns:
        dict: bloques ({'sg'|'rg'|'pp'|'todos': {filas, sin_centro_n,
        sin_centro_importe, total}}), referencia2 (Series for the helper
        column, None when there is no data), filas_sin_uuid and avisos.
    """
    if monivoi_sin_datos or monivoi.empty:
        return {
            "bloques":        _mtd_bloques_vacios(),
            "referencia2":    None,
            "filas_sin_uuid": 0,
            "avisos": [
                "MTD: el archivo de Monivoi vino sin datos; las cuatro matrices "
                "quedaron vacías."
            ],
            # No columns were resolved (we never got to read the file), so the
            # legend falls back to the canonical names. With no data there is
            # nothing to audit anyway.
            "columnas": {
                "uuid":   MTD_COL_UUID_CANDIDATAS[0],
                "centro": MTD_COL_CENTRO_CANDIDATAS[0],
            },
        }

    # Monivoi columns. They are resolved by name (accepting spelling variants)
    # and fail with a clear message if missing: stopping is better than
    # summing the wrong column.
    col_ref    = col_por_candidatos(monivoi, MTD_COL_REFERENCIA_CANDIDATAS,
                                     "referencia", origen="Monivoi")
    col_uuid   = col_por_candidatos(monivoi, MTD_COL_UUID_CANDIDATAS,
                                     "UUID", origen="Monivoi")
    col_centro = col_por_candidatos(monivoi, MTD_COL_CENTRO_CANDIDATAS,
                                     "centro de distribución", origen="Monivoi")
    col_valor  = col_por_candidatos(monivoi, MTD_COL_VALOR_CANDIDATAS,
                                     "valor neto", origen="Monivoi")

    avisos = []

    # Grouping key: first 2 characters of 'Referencia'. It is kept as is (for
    # column A of the sheet) and compared upper-cased.
    referencia2 = monivoi[col_ref].astype(str).str.strip().str[:2]
    ref2_norm   = referencia2.str.upper()
    centro_norm = monivoi[col_centro].astype(str).str.strip()

    con_uuid = ~monivoi[col_uuid].map(_uuid_vacio)
    filas_sin_uuid = int((~con_uuid).sum())
    if filas_sin_uuid:
        avisos.append(
            f"MTD: {filas_sin_uuid:,} fila(s) de Monivoi se excluyeron por traer "
            f"el UUID vacío o nulo."
        )

    # Universe of rows that go into SOME matrix. The amount is parsed only here
    # (the union of the three prefixes == the 'TODOS' matrix).
    mascara_union = con_uuid & ref2_norm.isin(MTD_PREFIJOS_TODOS)
    # astype(float) is NOT cosmetic: if the mask lets no row through, .map()
    # returns an empty text-typed Series (the Monivoi is read as text: 'object'
    # in pandas 2, 'str' in pandas 3), and with pandas 3 its .sum() gives ''
    # (the sum of zero strings) instead of 0.0, breaking the rounding. Pinning
    # the dtype makes the empty case 0.0 on any pandas version.
    valores = monivoi.loc[mascara_union, col_valor].map(
        partial(parse_importe_us, columna=col_valor)
    ).astype("float64")
    base = pd.DataFrame({
        "_ref2":   ref2_norm[mascara_union],
        "_centro": centro_norm[mascara_union],
        "_valor":  valores,
    })

    bloques = {}
    vacios  = []
    for bloque in MTD_BLOQUES:
        sub = base[base["_ref2"].isin(bloque["prefijos"])]

        sin_centro         = sub["_centro"] == ""
        sin_centro_n       = int(sin_centro.sum())
        sin_centro_importe = round(float(sub.loc[sin_centro, "_valor"].sum()), 2)

        con_centro = sub.loc[~sin_centro]
        if con_centro.empty:
            filas = []
        else:
            # groupby sorts the keys: centros come out alphabetically, as in 'TD'.
            serie = con_centro.groupby("_centro")["_valor"].sum()
            filas = [(centro, round(float(valor), 2)) for centro, valor in serie.items()]

        # The total adds the ALREADY ROUNDED values, so the column adds up to
        # exactly what is shown on screen.
        total = round(sum(valor for _, valor in filas) + sin_centro_importe, 2)

        if not filas and not sin_centro_n:
            vacios.append(bloque["titulo"])
        if sin_centro_n:
            avisos.append(
                f"MTD ({bloque['titulo']}): {sin_centro_n:,} fila(s) sin centro "
                f"(se muestran como '{MTD_ETIQUETA_SIN_CENTRO}')."
            )

        bloques[bloque["clave"]] = {
            "filas":              filas,
            "sin_centro_n":       sin_centro_n,
            "sin_centro_importe": sin_centro_importe,
            "total":              total,
        }

    if vacios:
        avisos.append(
            f"MTD: sin filas para {', '.join(vacios)} "
            f"(no hubo Referencia con ese prefijo y UUID)."
        )

    # Reconciliation check: 'TODOS' is the union of the other three, so its
    # total MUST be the sum of the three. If it does not match, warn instead of
    # silently delivering the file (the tolerance covers rounding to cents).
    suma_partes = round(
        bloques["sg"]["total"] + bloques["rg"]["total"] + bloques["pp"]["total"], 2
    )
    diferencia = round(bloques["todos"]["total"] - suma_partes, 2)
    if abs(diferencia) > MTD_TOLERANCIA_CUADRE:
        avisos.append(
            f"MTD: el total de 'TODOS' ({bloques['todos']['total']:,.2f}) no "
            f"cuadra con SG+RG+PP ({suma_partes:,.2f}); diferencia "
            f"{diferencia:,.2f}. Revisa el archivo antes de usar la matriz."
        )

    return {
        "bloques":        bloques,
        "referencia2":    referencia2,
        "filas_sin_uuid": filas_sin_uuid,
        "avisos":         avisos,
        # Names the columns REALLY have in this export (they accept spelling
        # variants: 'Centro'/'Centro de distribución', 'UUID'/'UUID CFDI'...).
        # The filter legend cites them, so the stakeholder reads the same
        # header the 'Monivoi' sheet shows.
        "columnas": {
            "uuid":   col_uuid,
            "centro": col_centro,
        },
    }


def insertar_referencia2(monivoi, mtd):
    """
    Insert the helper column 'Referencia 2' AS THE FIRST COLUMN (A) of the
    Monivoi sheet, so the stakeholder can audit the 'MTD' matrices by
    filtering by hand.

    Call it AFTER `preparar_matrices_mtd` (the matrices use their own series,
    not the sheet). It uses a free column name, so it never overwrites a
    same-named column the export already had; in that case a warning is
    appended to mtd["avisos"].

    It MUTATES `monivoi` (new first column) and, in that case, mtd["avisos"].
    It also records the FINAL name of the column in mtd["nombre_ref2"], which
    is what the 'MTD' filter legends cite. It does nothing when the Monivoi had
    no data (mtd["referencia2"] is None).
    """
    if mtd["referencia2"] is not None:
        nombre_ref2 = _nombre_columna_libre(monivoi, MTD_COL_REFERENCIA2)
        if nombre_ref2 != MTD_COL_REFERENCIA2:
            mtd["avisos"].append(
                f"MTD: la Monivoi ya traía una columna '{MTD_COL_REFERENCIA2}'; "
                f"la auxiliar se agregó como '{nombre_ref2}' para no pisarla."
            )
        monivoi.insert(0, nombre_ref2, mtd["referencia2"])
        # The filter legend of each matrix must point at the column the
        # stakeholder will REALLY find on the 'Monivoi' sheet. When the export
        # already brought a 'Referencia 2', the helper went in under another
        # name, and a legend still saying 'Referencia 2' would send the audit
        # to the wrong (pre-existing) column.
        mtd["nombre_ref2"] = nombre_ref2


def _leyenda_filtros(bloque, nombre_ref2, columnas):
    """
    Filter legend of an 'MTD' block, built FROM THAT BLOCK'S OWN constants
    (its prefixes), so the text can never drift from what the matrix really
    sums: adding a prefix to MTD_BLOQUES changes the legend too.

    It states the three conditions `preparar_matrices_mtd` applies — source
    sheet, 'Referencia 2' (first 2 characters of 'Referencia', upper-cased)
    and a non-empty UUID — plus what happens to the rows that pass the filter
    without a centro, which are the ones a manual check usually misses.

    `nombre_ref2` is the REAL name the helper column ended up with on the
    'Monivoi' sheet (see insertar_referencia2) and `columnas` carries the
    resolved names of the UUID and Centro columns of this export, so every
    column the legend names is one the stakeholder can actually filter.

    Each line is kept short on purpose: the first block sits in column A and
    only has columns B and C to overflow into before running into the next
    block's legend.
    """
    prefijos = " o ".join(f"'{p}'" for p in sorted(bloque["prefijos"]))
    return [
        f"Filtros — hoja '{HOJA_MONIVOI}':",
        f"{nombre_ref2} = {prefijos}; {columnas['uuid']} con dato",
        f"{columnas['centro']} vacío → renglón '{MTD_ETIQUETA_SIN_CENTRO}'",
    ]


def escribir_hoja_mtd(ws, mtd):
    """
    Write the four step-16 matrices onto the 'MTD' sheet.

    Above each matrix (rows 1-3 of its own column) goes the legend of the
    filters that built it, so the stakeholder can reproduce any figure by hand
    from the 'Monivoi' sheet. Call it AFTER `insertar_referencia2`, so the
    legend can cite the helper column by its final name.
    """
    # Name the helper column really ended up with (it changes only when the
    # export already had one with that name). Without data there is no helper
    # column and the matrices are empty, so the default name is fine.
    nombre_ref2 = mtd.get("nombre_ref2", MTD_COL_REFERENCIA2)
    columnas = mtd["columnas"]

    for bloque in MTD_BLOQUES:
        datos     = mtd["bloques"][bloque["clave"]]
        col_centro = bloque["col_centro"]
        col_valor  = bloque["col_valor"]

        # Filter legend (rows 1-3, in the centro column of this block).
        escribir_filtros(ws, MTD_FILA_FILTROS, col_centro,
                         _leyenda_filtros(bloque, nombre_ref2, columnas))

        # Title (row 4) and headers (row 5).
        ws.cell(row=MTD_FILA_TITULO, column=col_centro,
                value=bloque["titulo"]).font = Font(bold=True)
        ws.cell(row=MTD_FILA_ENCABEZADO, column=col_centro,
                value=MTD_ENCABEZADO_CENTRO).font = Font(bold=True)
        ws.cell(row=MTD_FILA_ENCABEZADO, column=col_valor,
                value=MTD_ENCABEZADO_VALOR).font = Font(bold=True)

        # Body: one row per centro.
        fila = MTD_FILA_INICIO
        for centro, valor in datos["filas"]:
            ws.cell(row=fila, column=col_centro, value=centro)
            ws.cell(row=fila, column=col_valor,
                    value=valor).number_format = TD_FORMATO_NUMERO
            fila += 1

        # Rows that passed the filter but had no centro: shown apart, so the
        # 'Grand Total' still reconciles with the filter.
        if datos["sin_centro_n"]:
            ws.cell(row=fila, column=col_centro,
                    value=MTD_ETIQUETA_SIN_CENTRO).font = Font(italic=True)
            celda = ws.cell(row=fila, column=col_valor,
                            value=datos["sin_centro_importe"])
            celda.number_format = TD_FORMATO_NUMERO
            celda.font = Font(italic=True)
            fila += 1

        # 'Grand Total' only if there was at least one data row.
        if fila > MTD_FILA_INICIO:
            ws.cell(row=fila, column=col_centro,
                    value=MTD_ETIQUETA_TOTAL).font = Font(bold=True)
            celda_total = ws.cell(row=fila, column=col_valor, value=datos["total"])
            celda_total.number_format = TD_FORMATO_NUMERO
            celda_total.font = Font(bold=True)

        # Comfortable widths to read centro and amount.
        ws.column_dimensions[get_column_letter(col_centro)].width = 12
        ws.column_dimensions[get_column_letter(col_valor)].width = 20


# ----------------------------------------------------------------------
# TD × MTD cross — live formulas and controls
# ----------------------------------------------------------------------
def _ref_hoja(nombre):
    """Sheet name as a formula reference: bare when allowed, quoted otherwise."""
    if (_RE_HOJA_SIN_COMILLAS.fullmatch(nombre)
            and nombre.upper() not in _HOJAS_RESERVADAS):
        return nombre
    return "'" + nombre.replace("'", "''") + "'"


def _texto_formula(texto):
    """Text literal for a formula: 'Sin timbre' -> "Sin timbre"."""
    return '"' + str(texto).replace('"', '""') + '"'


def _formulas_cruce(cruce, fila):
    """
    The three formulas of ONE row of the TD × MTD cross, in this order:
    (Factura Monitor, Diferencia, Observación). Example (Facturas, row 6):

        E6: =IFERROR(XLOOKUP(A6,MTD!A:A,MTD!B:B),"Sin timbre")
        F6: =IFERROR(D6+E6,"Sin timbre")
        G6: =IF(F6="Sin timbre","Sin timbre",IF(ROUND(F6,2)<-0.01,"","Sin diferencia"))

    In the file XLOOKUP is written as `_xlfn.XLOOKUP`: functions newer than
    Excel 2007 need that prefix, otherwise Excel shows #NAME?.
    """
    if cruce["operador"] not in ("+", "-"):
        raise ValueError(
            f"Cruce TD ({cruce['etiqueta']}): operador de diferencia inválido "
            f"'{cruce['operador']}' (se esperaba '+' o '-')."
        )
    if cruce["revisar_si"] not in ("<", ">"):
        raise ValueError(
            f"Cruce TD ({cruce['etiqueta']}): comparación inválida "
            f"'{cruce['revisar_si']}' (se esperaba '<' o '>')."
        )

    letra      = get_column_letter
    bloque_mtd = _MTD_POR_CLAVE[cruce["mtd_clave"]]
    hoja_mtd   = _ref_hoja(HOJA_MTD)
    col_centro = letra(bloque_mtd["col_centro"])
    col_valor  = letra(bloque_mtd["col_valor"])

    cedis   = f"{letra(cruce['col_cedis'])}{fila}"
    contab  = f"{letra(cruce['col_contab'])}{fila}"
    monitor = f"{letra(cruce['col_monitor'])}{fila}"
    dif     = f"{letra(cruce['col_diferencia'])}{fila}"

    sin_timbre     = _texto_formula(TD_TEXTO_SIN_TIMBRE)
    sin_diferencia = _texto_formula(TD_TEXTO_SIN_DIFERENCIA)
    umbral = (-TD_CRUCE_TOLERANCIA if cruce["revisar_si"] == "<"
              else TD_CRUCE_TOLERANCIA)

    f_monitor = (
        f"=IFERROR(_xlfn.XLOOKUP({cedis},{hoja_mtd}!{col_centro}:{col_centro},"
        f"{hoja_mtd}!{col_valor}:{col_valor}),{sin_timbre})"
    )
    f_diferencia = f"=IFERROR({contab}{cruce['operador']}{monitor},{sin_timbre})"
    f_observacion = (
        f"=IF({dif}={sin_timbre},{sin_timbre},"
        f"IF(ROUND({dif},2){cruce['revisar_si']}{umbral:g},\"\",{sin_diferencia}))"
    )
    return f_monitor, f_diferencia, f_observacion


def escribir_cruce_td_mtd(ws_td, td_resumen):
    """
    Write the cross against 'MTD' onto the 'TD' sheet: three new columns per
    block (see TD_CRUCES), with the header in row 5 and a formula in every
    centro row PLUS the 'Grand Total' row.

    Headers are always written (like the rest of 'TD'); if the block has no
    centros, it has no total row either and no formula is written. Amount
    cells get the thousands/2-decimals format ("Sin timbre" looks the same,
    because it is text).
    """
    negrita = Font(bold=True)
    for cruce in TD_CRUCES:
        columnas = (
            (cruce["col_monitor"],     TD_ENCABEZADO_MONITOR),
            (cruce["col_diferencia"],  TD_ENCABEZADO_DIFERENCIA),
            (cruce["col_observacion"], TD_ENCABEZADO_OBSERVACION),
        )
        for col, encabezado in columnas:
            ws_td.cell(row=TD_FILA_ENCABEZADO, column=col, value=encabezado).font = negrita
            ws_td.column_dimensions[get_column_letter(col)].width = max(16, len(encabezado) + 3)

        n_cedis = len(td_resumen[cruce["td_clave"]]["cedis"])
        if n_cedis == 0:
            continue
        fila_total = TD_FILA_INICIO + n_cedis   # same formula as escribir_matriz_td (Gallo)

        for fila in range(TD_FILA_INICIO, fila_total + 1):
            f_monitor, f_diferencia, f_observacion = _formulas_cruce(cruce, fila)
            celda_monitor = ws_td.cell(row=fila, column=cruce["col_monitor"], value=f_monitor)
            celda_dif     = ws_td.cell(row=fila, column=cruce["col_diferencia"], value=f_diferencia)
            celda_obs     = ws_td.cell(row=fila, column=cruce["col_observacion"], value=f_observacion)
            celda_monitor.number_format = TD_FORMATO_NUMERO
            celda_dif.number_format     = TD_FORMATO_NUMERO
            if fila == fila_total:
                for celda in (celda_monitor, celda_dif, celda_obs):
                    celda.font = negrita


def avisos_cruce_td_mtd(td_resumen, mtd):
    """
    Cross checks that the formulas do NOT show on their own.

    The XLOOKUP goes from TD to MTD: a centro that the monitor did stamp but
    that has no row in TD never shows up in a row of its own; it only weighs
    in the 'Grand Total'. Those centros are listed here so the stakeholder
    knows WHAT explains the difference in the total.

    Matching works like XLOOKUP: case-insensitive and keeping the FIRST match.
    That also uncovers the rare case of centros that only differ in case
    ('80AB' / '80ab'): in MTD the second one would stay unmatched, and in TD
    both rows would get the same monitor amount.

    Returns:
        list[str]: warnings ready for the GUI summary.
    """
    avisos = []
    for cruce in TD_CRUCES:
        etiqueta    = f"Cruce TD ({cruce['etiqueta']})"
        cedis_td    = td_resumen[cruce["td_clave"]]["cedis"]
        centros_mtd = [centro for centro, _ in mtd["bloques"][cruce["mtd_clave"]]["filas"]]

        # MTD row that XLOOKUP would find for each key (the first one).
        primero_mtd = {}
        for centro in centros_mtd:
            primero_mtd.setdefault(centro.casefold(), centro)
        claves_td  = {centro.casefold() for centro in cedis_td}
        alcanzados = {primero_mtd[clave] for clave in claves_td if clave in primero_mtd}
        sin_renglon = [centro for centro in centros_mtd if centro not in alcanzados]

        if sin_renglon:
            listado = ", ".join(sin_renglon[:5]) + ("..." if len(sin_renglon) > 5 else "")
            if cedis_td:
                avisos.append(
                    f"{etiqueta}: {len(sin_renglon):,} centro(s) del monitor sin "
                    f"renglón en TD ({listado}); su importe sólo pesa en el "
                    f"'{ETIQUETA_GRAND_TOTAL}'."
                )
            else:
                avisos.append(
                    f"{etiqueta}: TD no trae centros en este bloque, pero el "
                    f"monitor sí ({listado}); no hay contra qué cruzar."
                )

        grupos_td = {}
        for centro in cedis_td:
            grupos_td.setdefault(centro.casefold(), []).append(centro)
        choques = [" / ".join(grupo) for grupo in grupos_td.values() if len(grupo) > 1]
        if choques:
            avisos.append(
                f"{etiqueta}: en TD hay centros que sólo cambian en mayúsculas/"
                f"minúsculas ({'; '.join(choques)}); XLOOKUP les asigna el mismo "
                f"importe del monitor."
            )
    return avisos
